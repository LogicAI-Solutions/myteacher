"""Relatório do aluno: o Word (download e anexo do e-mail) e o resumo do mês para os pais."""
import base64
import io
from typing import Optional

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

MONTHS_PT = ["janeiro", "fevereiro", "março", "abril", "maio", "junho",
             "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


def build_student_report_docx(stats: dict, month: Optional[int] = None, year: Optional[int] = None,
                              chart_image: Optional[str] = None) -> bytes:
    document = Document()
    
    # Title
    title_text = 'Relatório de Desempenho'
    if month and year:
        title_text += f' - {month:02d}/{year}'
    
    title = document.add_heading(title_text, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Student Info Section
    document.add_heading('Informações do Aluno', level=1)
    p = document.add_paragraph()
    p.add_run('Nome: ').bold = True
    p.add_run(f'{stats["student"].name}\n')
    if stats["student"].parent_name:
        p.add_run('Responsável: ').bold = True
        p.add_run(f'{stats["student"].parent_name}\n')
    
    # Statistics Section
    document.add_heading('Resumo de Atividades', level=1)
    table_stats = document.add_table(rows=1, cols=4)
    table_stats.style = 'Table Grid'
    hdr_stats = table_stats.rows[0].cells
    hdr_stats[0].text = 'Total Aulas'
    hdr_stats[1].text = 'Presenças'
    hdr_stats[2].text = 'Frequência'
    hdr_stats[3].text = 'Média Notas'
    
    # Style Headers
    for cell in hdr_stats:
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        
    row_stats = table_stats.add_row().cells
    row_stats[0].text = str(stats["total_classes"])
    row_stats[1].text = str(stats["present"])
    row_stats[2].text = f'{stats["attendance_rate"]}%'
    row_stats[3].text = f'{stats["avg_grade"]}'

    # Evolution Chart Section
    if chart_image:
        document.add_heading('Gráfico de Evolução', level=1)
        try:
            # Decode base64 image
            # Format usually: "data:image/png;base64,....."
            if "," in chart_image:
                header, encoded = chart_image.split(",", 1)
            else:
                encoded = chart_image
            
            image_data = base64.b64decode(encoded)
            image_stream = io.BytesIO(image_data)
            
            # Add picture centered
            document.add_picture(image_stream, width=Inches(6))
            last_paragraph = document.paragraphs[-1] 
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            print(f"Error adding image: {e}")
            document.add_paragraph("[Erro ao incluir o gráfico]")

    
    # Detailed History
    document.add_heading('Histórico Detalhado', level=1)
    
    table = document.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Data'
    hdr_cells[1].text = 'Conteúdo/Descrição'
    hdr_cells[2].text = 'Status'
    hdr_cells[3].text = 'Nota'
    hdr_cells[4].text = 'Observação'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].bold = True
    
    for log in stats["logs"]:
        row_cells = table.add_row().cells
        row_cells[0].text = log.session.date.strftime('%d/%m/%Y')
        row_cells[1].text = str(log.session.description)
        
        # Translate status
        # Translate status
        status_map = {
            'PRESENT': 'Presente', 
            'ABSENT': 'Ausente', 
            'LATE': 'Atrasado', 
            'Justified': 'Justificado',
            'present': 'Presente',
            'absent': 'Ausente',
            'late': 'Atrasado',
            'justified': 'Justificado'
        }
        row_cells[2].text = status_map.get(log.status, log.status)
        
        row_cells[3].text = str(log.grade) if log.grade is not None else '-'
        row_cells[4].text = str(log.observation) if log.observation else '-'
        
    file_stream = io.BytesIO()
    document.save(file_stream)
    return file_stream.getvalue()


def student_report_filename(stats: dict, month: Optional[int] = None, year: Optional[int] = None) -> str:
    # Filename construction
    safe_name = stats['student'].name.replace(' ', '_')
    filename = f"Relatorio_{safe_name}"
    
    if month and year:
        filename += f"_{month:02d}_{year}"
    elif year:
        filename += f"_{year}"
        
    filename += ".docx"
    return filename


def _fmt_grade(value: float) -> str:
    return f"{value:.1f}".rstrip("0").rstrip(".").replace(".", ",")


def monthly_summary(stats: dict, month: int, year: int, sender: str) -> dict:
    """Resumo do mês para os pais: texto (WhatsApp e e-mail) e HTML (e-mail)."""
    student = stats["student"]
    logs = stats["logs"]
    month_name = f"{MONTHS_PT[month - 1]} de {year}"
    total, present = stats["total_classes"], stats["present"]
    grades = [l.grade for l in logs if l.grade is not None]
    essays = len([l for l in logs if getattr(l, "essay_delivered", False)])
    notes = [(l.session.date, l.observation.strip()) for l in logs if l.observation and l.observation.strip()]

    lines = [f"Resumo de {month_name} — {student.name}", ""]
    if total:
        lines.append(f"Frequência: {present} de {total} aula{'s' if total != 1 else ''} ({round(stats['attendance_rate'])}%)")
    else:
        lines.append("Não houve aulas registradas neste mês.")
    if grades:
        lines.append(f"Média das notas: {_fmt_grade(sum(grades) / len(grades))} ({len(grades)} avaliação{'ões' if len(grades) != 1 else ''})")
    if essays:
        lines.append(f"Redações entregues: {essays}")
    if notes:
        lines.append("")
        lines.append("Observações do professor:")
        lines += [f"• {d.strftime('%d/%m')}: {text}" for d, text in notes]
    lines += ["", f"— {sender}"]
    text = "\n".join(lines)

    import html as _html
    html = (
        f"<p>Olá! Segue o resumo de <strong>{_html.escape(month_name)}</strong> de "
        f"<strong>{_html.escape(student.name)}</strong>.</p>"
        f"<pre style=\"font-family:Arial,sans-serif;white-space:pre-wrap;font-size:14px\">{_html.escape(chr(10).join(lines[2:-2]))}</pre>"
        f"<p>O relatório completo está em anexo.</p><p>{_html.escape(sender)}</p>"
    )
    return {"text": text, "html": html, "subject": f"Resumo de {month_name} — {student.name}", "has_classes": bool(total)}
